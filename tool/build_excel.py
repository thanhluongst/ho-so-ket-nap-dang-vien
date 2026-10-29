# -*- coding: utf-8 -*-
"""Tạo file Excel mẫu để nhập thông tin hồ sơ rồi import vào tool (bản máy và bản online)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from ket_nap import MAU_NX, TABS

HERE = Path(__file__).parent
ROOT = HERE.parent
OUT = ROOT / 'Mau nhap thong tin.xlsx'

DO = 'A3221C'
GHI_CHU = {'date': 'Ngày/tháng/năm, VD: 05/03/1998', 'int': 'Số',
           'long': 'Chọn mẫu trong danh sách hoặc tự viết. Alt+Enter để xuống dòng'}

wb = Workbook()
ws = wb.active
ws.title = 'Thông tin'
ds = wb.create_sheet('DanhSach')   # tên các mẫu nhận xét cho ô chọn
ds.sheet_state = 'hidden'

thin = Side(style='thin', color='D9D2D0')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
ws.column_dimensions['A'].width = 40
ws.column_dimensions['B'].width = 60
ws.column_dimensions['C'].width = 38
ws.column_dimensions['D'].width = 14

ws['A1'] = 'THÔNG TIN HỒ SƠ KẾT NẠP ĐẢNG VIÊN'
ws['A1'].font = Font(bold=True, size=14, color=DO)
ws.merge_cells('A1:D1')
ws['A2'] = ('Nhập vào cột "Nội dung nhập", lưu file rồi bấm "Nhập từ Excel" trên tool. '
            'Ô để trống sẽ để trống trên mẫu. Không sửa cột "Mã".')
ws['A2'].alignment = Alignment(wrap_text=True)
ws.merge_cells('A2:D2')
ws.row_dimensions[2].height = 32

for col, title in enumerate(['Thông tin', 'Nội dung nhập', 'Ghi chú', 'Mã'], 1):
    c = ws.cell(row=4, column=col, value=title)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor=DO)
    c.border = border
ws.freeze_panes = 'A5'

row = 5
for col, key in enumerate(['nhan_thuc', 'ly_lich', 'uu_diem', 'khuyet_diem'], 1):
    for i, m in enumerate(MAU_NX[key], 1):
        ds.cell(row=i, column=col, value=m['ten'])

for tab_name, fields in TABS:
    c = ws.cell(row=row, column=1, value=tab_name.upper())
    c.font = Font(bold=True, color=DO)
    c.fill = PatternFill('solid', fgColor='F6E7E5')
    for col in range(1, 5):
        ws.cell(row=row, column=col).fill = PatternFill('solid', fgColor='F6E7E5')
    row += 1
    for key, label, kind, default, note in fields:
        ws.cell(row=row, column=1, value=label).alignment = Alignment(vertical='top', wrap_text=True)
        val = MAU_NX[key][0]['ten'] if kind == 'long' else default
        v = ws.cell(row=row, column=2, value=val or None)
        v.number_format = '@'   # giữ nguyên chữ, Excel không tự đổi ngày/số
        v.alignment = Alignment(vertical='top', wrap_text=True)
        n = ws.cell(row=row, column=3, value=' · '.join(x for x in (GHI_CHU.get(kind, ''), note) if x) or None)
        n.font = Font(color='6D6361', size=10)
        n.alignment = Alignment(vertical='top', wrap_text=True)
        k = ws.cell(row=row, column=4, value=key)
        k.font = Font(color='A89E9B', size=9)
        for col in range(1, 5):
            ws.cell(row=row, column=col).border = border
        if kind == 'long':
            ws.row_dimensions[row].height = 60
            colx = 'ABCD'[['nhan_thuc', 'ly_lich', 'uu_diem', 'khuyet_diem'].index(key)]
            dv = DataValidation(type='list', formula1=f'=DanhSach!${colx}$1:${colx}${len(MAU_NX[key])}',
                                allow_blank=True, showErrorMessage=False)
            ws.add_data_validation(dv)
            dv.add(v)
        row += 1
    row += 1

wb.save(OUT)
print('Đã tạo', OUT, '(chạy build_web.py để đưa vào bản online)')
