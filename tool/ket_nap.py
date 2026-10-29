# -*- coding: utf-8 -*-
"""Tool lập hồ sơ kết nạp đảng viên: nhập thông tin một lần, xuất Mẫu 1, 3, 5, 6, 8-KNĐ."""
import datetime
import json
import os
import re
import sys
from pathlib import Path

from docxtpl import DocxTemplate

HERE = Path(__file__).parent
MAU_DIR = HERE / 'mau'
MAU_DO_DIR = HERE / 'mau_do'   # bản tô đỏ chữ được điền
HOSO_DIR = HERE.parent / 'HoSo'
DON_VI_FILE = HERE / 'don_vi.json'

TRONG = '……………'        # chỗ trống khi chưa nhập, để có thể viết tay
TRONG_NGAY = '……'

# Mẫu nhận xét gợi ý (sửa/thêm trong mau_nhan_xet.json); mẫu đầu tiên là mặc định
MAU_NX = json.loads((HERE / 'mau_nhan_xet.json').read_text(encoding='utf-8'))
NX_KEYS = ['nhan_thuc', 'ly_lich', 'uu_diem', 'khuyet_diem']
NHAN_THUC, LY_LICH, UU_DIEM, KHUYET_DIEM = (MAU_NX[k][0]['noi_dung'] for k in NX_KEYS)


def thay_the(text, v):
    """Thay {ten}, {nghe}, {chuc_vu}, {don_vi}, {chi_bo} bằng thông tin đã nhập (bỏ qua ô chưa nhập)."""
    vals = {'ten': v.get('ho_ten', ''), 'nghe': v.get('nghe_nghiep', '').lower(),
            'chuc_vu': v.get('chuc_vu', '').lower(), 'don_vi': v.get('don_vi', ''), 'chi_bo': v.get('chi_bo', '')}
    return re.sub(r'\{(\w+)\}', lambda m: vals.get(m.group(1)) or m.group(0), text)


# (khóa, nhãn, kiểu, mặc định, ghi chú). Kiểu: text | date | int | long.
# Các khóa nằm trong tab "Đơn vị" được lưu lại cho lần sau (don_vi.json).
TABS = [
    ('Đơn vị', [
        ('truong', 'Tên trường', 'text', 'THCS Nhơn Mỹ', 'Không ghi chữ "Trường"'),
        ('chi_bo', 'Tên chi bộ', 'text', '', 'VD: Giáo viên 1 (không ghi chữ "Chi bộ")'),
        ('dia_danh', 'Địa danh ghi ngày tháng', 'text', 'Nhơn Mỹ', ''),
        ('bt_chi_bo', 'Họ tên Bí thư chi bộ', 'text', '', ''),
        ('bt_dang_uy', 'Họ tên Bí thư Đảng ủy', 'text', '', ''),
        ('dv_ct', 'Số đảng viên chính thức của chi bộ', 'int', '', ''),
        ('dv_db', 'Số đảng viên dự bị của chi bộ', 'int', '0', ''),
        ('uv_tong', 'Tổng số ủy viên BCH Đảng bộ', 'int', '', ''),
    ]),
    ('Quần chúng', [
        ('ho_ten', 'Họ và tên', 'text', '', 'Bắt buộc'),
        ('ngay_sinh', 'Ngày sinh', 'date', '', 'dd/mm/yyyy'),
        ('noi_sinh', 'Nơi sinh', 'text', '', ''),
        ('que_quan', 'Quê quán', 'text', '', ''),
        ('dan_toc', 'Dân tộc', 'text', 'Kinh', ''),
        ('ton_giao', 'Tôn giáo', 'text', 'Không', ''),
        ('hoc_van', 'Trình độ học vấn', 'text', '12/12', ''),
        ('cu_tru', 'Nơi cư trú', 'text', '', ''),
        ('nghe_nghiep', 'Nghề nghiệp', 'text', 'Giáo viên', ''),
        ('don_vi', 'Đơn vị công tác', 'text', 'Trường THCS Nhơn Mỹ', ''),
        ('chuc_vu', 'Chức vụ chính quyền, đoàn thể', 'text', '', ''),
        ('ngay_doan', 'Ngày vào Đoàn', 'date', '', ''),
        ('doan_noi', 'Nơi vào Đoàn', 'text', '', ''),
        ('ngay_ctd', 'Ngày được xét cảm tình Đảng', 'date', '', ''),
        ('ctd_chi_bo', 'Tại chi bộ', 'text', '', 'Để trống = chi bộ hiện tại'),
        ('ngay_don', 'Ngày làm đơn (Mẫu 1)', 'date', '', ''),
    ]),
    ('Người giới thiệu', [
        ('gt_ten', 'Họ tên đảng viên giới thiệu', 'text', '', ''),
        ('gt_ns', 'Ngày sinh', 'date', '', ''),
        ('gt_vd', 'Ngày vào Đảng', 'date', '', ''),
        ('gt_ct', 'Ngày chính thức', 'date', '', ''),
        ('gt_cv_dang', 'Chức vụ trong Đảng', 'text', 'Đảng viên', ''),
        ('gt_cv_cq', 'Chức vụ chính quyền', 'text', 'Giáo viên', ''),
        ('gt_que', 'Quê quán', 'text', '', ''),
        ('gt_chi_bo', 'Đang sinh hoạt tại chi bộ', 'text', '', 'Để trống = chi bộ hiện tại'),
        ('ngay_pc', 'Ngày được phân công giúp đỡ', 'date', '', ''),
        ('ngay_gt_ky', 'Ngày ký giấy giới thiệu (Mẫu 3)', 'date', '', ''),
    ]),
    ('Đoàn thể & Cư trú', [
        ('ngay_m5', 'Ngày lập tổng hợp (Mẫu 5)', 'date', '', ''),
        ('ten_doan_the', 'Tên các đoàn thể nơi sinh hoạt', 'text', '', 'VD: Công đoàn, Chi đoàn giáo viên Trường ...'),
        ('so_dt', 'Tổng số đồng chí đại diện đoàn thể được hỏi', 'int', '', ''),
        ('dt_tan', 'Số tán thành (đoàn thể)', 'int', '', 'Để trống = tất cả tán thành'),
        ('dt_lydo', 'Lý do không tán thành (đoàn thể)', 'text', '', ''),
        ('ten_ct', 'Tên chi ủy nơi cư trú', 'text', '', 'VD: Chi bộ ấp ...'),
        ('so_ct', 'Số đồng chí chi ủy nơi cư trú được hỏi', 'int', '', ''),
        ('ct_tan', 'Số tán thành (nơi cư trú)', 'int', '', 'Để trống = tất cả tán thành'),
        ('ct_lydo', 'Lý do không tán thành (nơi cư trú)', 'text', '', ''),
    ]),
    ('Họp chi bộ', [
        ('so_nq_cb', 'Số Nghị quyết chi bộ', 'text', '', 'VD: 05'),
        ('ngay_cb', 'Ngày họp / ra nghị quyết', 'date', '', ''),
        ('vm_ct', 'Số đảng viên chính thức vắng', 'int', '0', ''),
        ('vm_db', 'Số đảng viên dự bị vắng', 'int', '0', ''),
        ('cb_lydo_vang', 'Lý do vắng mặt', 'text', '', ''),
        ('cb_chu_tri', 'Chủ trì hội nghị', 'text', '', 'Để trống = Bí thư chi bộ'),
        ('cb_chuc_vu', 'Chức vụ người chủ trì', 'text', 'Bí thư chi bộ', ''),
        ('cb_thu_ky', 'Thư ký hội nghị', 'text', '', ''),
        ('cb_tan', 'Số đảng viên chính thức tán thành', 'int', '', 'Để trống = tất cả chính thức có mặt'),
        ('cb_lydo_khong', 'Lý do không tán thành', 'text', '', ''),
    ]),
    ('Họp Đảng ủy', [
        ('so_nq_du', 'Số Nghị quyết Đảng ủy', 'text', '', ''),
        ('ngay_du', 'Ngày họp / ra nghị quyết', 'date', '', ''),
        ('uv_vm', 'Số ủy viên vắng mặt', 'int', '0', ''),
        ('du_lydo_vang', 'Lý do vắng mặt', 'text', '', ''),
        ('du_chu_tri', 'Chủ trì hội nghị', 'text', '', 'Để trống = Bí thư Đảng ủy'),
        ('du_chuc_vu', 'Chức vụ người chủ trì', 'text', 'Bí thư Đảng ủy', ''),
        ('du_thu_ky', 'Thư ký hội nghị', 'text', '', ''),
        ('du_dy', 'Số ủy viên đồng ý', 'int', '', 'Để trống = tất cả có mặt'),
        ('du_lydo_khong', 'Lý do không đồng ý', 'text', '', ''),
    ]),
    ('Nhận xét', [
        ('nhan_thuc', 'Nhận thức về Đảng (Mẫu 1, người làm đơn tự viết)', 'long', NHAN_THUC, ''),
        ('ly_lich', 'Về lý lịch (Mẫu 3, 6, 8)', 'long', LY_LICH, ''),
        ('uu_diem', 'Ưu điểm (Mẫu 3, 5, 6, 8)', 'long', UU_DIEM, ''),
        ('khuyet_diem', 'Khuyết điểm (Mẫu 3, 5, 6, 8)', 'long', KHUYET_DIEM, ''),
    ]),
]
FIELDS = {f[0]: f for _, fs in TABS for f in fs}
DON_VI_KEYS = [f[0] for f in TABS[0][1]]

MAU = [
    ('mau1', 'Mau 1-KND Don xin vao Dang'),
    ('mau3', 'Mau 3-KND Giay gioi thieu'),
    ('mau5', 'Mau 5-KND Tong hop y kien doan the va noi cu tru'),
    ('mau6', 'Mau 6-KND Nghi quyet chi bo'),
    ('mau8', 'Mau 8-KND Nghi quyet Dang uy'),
]


class LoiNhap(Exception):
    pass


def defaults():
    return {k: f[3] for k, f in FIELDS.items()}


def parse_date(key, s):
    s = s.strip()
    if not s:
        return None
    m = re.fullmatch(r'(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{4})', s)
    if not m:
        raise LoiNhap(f'"{FIELDS[key][1]}": ngày phải có dạng dd/mm/yyyy (đang là "{s}")')
    d, mo, y = map(int, m.groups())
    try:
        datetime.date(y, mo, d)
    except ValueError:
        raise LoiNhap(f'"{FIELDS[key][1]}": ngày {s} không hợp lệ')
    return f'{d:02d}', f'{mo:02d}', str(y)


def parse_int(key, s):
    s = str(s).strip()
    if not s:
        return None
    if not s.isdigit():
        raise LoiNhap(f'"{FIELDS[key][1]}" phải là số (đang là "{s}")')
    return int(s)


def pct(a, b):
    if a is None or not b:
        return TRONG_NGAY
    v = round(a * 100 / b, 2)
    return (str(int(v)) if v == int(v) else f'{v:g}').replace('.', ',')


def lines(s):
    return [x.strip() for x in s.splitlines() if x.strip()]


def build_context(data):
    """Từ dữ liệu nhập → toàn bộ trường cần điền vào các mẫu."""
    v = {k: str(data.get(k, '')).strip() for k in FIELDS}
    if not v['ho_ten']:
        raise LoiNhap('Chưa nhập họ tên quần chúng')
    n = {k: parse_int(k, v[k]) for k, f in FIELDS.items() if f[2] == 'int'}

    ctx = {}
    for k, f in FIELDS.items():
        if f[2] == 'text':
            ctx[k] = v[k] or TRONG
    ctx['ctd_chi_bo'] = v['ctd_chi_bo'] or v['chi_bo'] or TRONG
    ctx['gt_chi_bo'] = v['gt_chi_bo'] or v['chi_bo'] or TRONG
    ctx['don_vi'] = v['don_vi'] or (f"Trường {v['truong']}" if v['truong'] else TRONG)
    ctx['cb_chu_tri'] = v['cb_chu_tri'] or v['bt_chi_bo'] or TRONG
    ctx['du_chu_tri'] = v['du_chu_tri'] or v['bt_dang_uy'] or TRONG
    ctx['TRUONG'] = v['truong'].upper() or TRONG
    ctx['CHI_BO'] = v['chi_bo'].upper() or TRONG
    ctx['DANG_UY'] = f"TRƯỜNG {v['truong'].upper()}" if v['truong'] else TRONG
    ctx['dang_bo'] = f"Trường {v['truong']}" if v['truong'] else TRONG
    ctx['so_nq_cb'] = v['so_nq_cb'] or '      '
    ctx['so_nq_du'] = v['so_nq_du'] or '      '

    # ngày tháng: <tiền tố>_d/_m/_y
    for key, prefix in [('ngay_sinh', 'ns'), ('ngay_doan', 'doan'), ('ngay_ctd', 'ctd'),
                        ('ngay_don', 'don'), ('gt_ns', 'gt_ns'), ('gt_vd', 'gt_vd'),
                        ('gt_ct', 'gt_ct'), ('ngay_pc', 'pc'), ('ngay_gt_ky', 'gt_ky'),
                        ('ngay_m5', 'm5'), ('ngay_cb', 'cb'), ('ngay_du', 'du')]:
        parts = parse_date(key, v[key]) or (TRONG_NGAY,) * 3
        ctx[f'{prefix}_d'], ctx[f'{prefix}_m'], ctx[f'{prefix}_y'] = parts

    def so(x):
        return TRONG_NGAY if x is None else str(x)

    def ly_do(khong, text):
        if khong == 0:
            return 'không.'
        return (text.rstrip('.') + '.') if text else TRONG

    def check(tan, tong, ten):
        if tan is not None and tong is not None and tan > tong:
            raise LoiNhap(f'{ten}: số tán thành ({tan}) lớn hơn tổng số ({tong})')

    # Mẫu 5
    for p, tong_key in (('dt', 'so_dt'), ('ct', 'so_ct')):
        tong = n[tong_key]
        tan = n[f'{p}_tan'] if n[f'{p}_tan'] is not None else tong
        check(tan, tong, 'Mẫu 5')
        khong = None if tan is None or tong is None else tong - tan
        ctx.update({f'{p}_tan': so(tan), f'{p}_tong': so(tong), f'{p}_pt': pct(tan, tong),
                    f'{p}_khong': so(khong), f'{p}_kpt': pct(khong, tong),
                    f'{p}_lydo': ly_do(khong, v[f'{p}_lydo'])})
    ctx['so_dt'], ctx['so_ct'] = so(n['so_dt']), so(n['so_ct'])

    # Mẫu 6
    dv_ct, dv_db = n['dv_ct'], n['dv_db'] or 0
    vm_ct, vm_db = n['vm_ct'] or 0, n['vm_db'] or 0
    if dv_ct is not None and (vm_ct > dv_ct or vm_db > dv_db):
        raise LoiNhap('Họp chi bộ: số vắng mặt lớn hơn số đảng viên')
    cm_ct = None if dv_ct is None else dv_ct - vm_ct
    cb_tan = n['cb_tan'] if n['cb_tan'] is not None else cm_ct
    check(cb_tan, cm_ct, 'Họp chi bộ (so với số chính thức có mặt)')
    cb_khong = None if cb_tan is None or cm_ct is None else cm_ct - cb_tan
    ctx.update({
        'tong_dv': so(None if dv_ct is None else dv_ct + dv_db), 'dv_ct': so(dv_ct), 'dv_db': so(dv_db),
        'cm_tong': so(None if cm_ct is None else cm_ct + dv_db - vm_db), 'cm_ct': so(cm_ct),
        'cm_db': so(dv_db - vm_db), 'vm_tong': so(vm_ct + vm_db), 'vm_ct': so(vm_ct), 'vm_db': so(vm_db),
        'cb_lydo_vang': v['cb_lydo_vang'] or ('không.' if vm_ct + vm_db == 0 else TRONG),
        'cb_tan': so(cb_tan), 'cb_pt': pct(cb_tan, dv_ct),
        'cb_khong': so(cb_khong), 'cb_kpt': pct(cb_khong, dv_ct),
        'cb_lydo_khong': ly_do(cb_khong, v['cb_lydo_khong']),
    })

    # Mẫu 8
    uv, uv_vm = n['uv_tong'], n['uv_vm'] or 0
    if uv is not None and uv_vm > uv:
        raise LoiNhap('Họp Đảng ủy: số vắng mặt lớn hơn tổng số ủy viên')
    uv_cm = None if uv is None else uv - uv_vm
    du_dy = n['du_dy'] if n['du_dy'] is not None else uv_cm
    check(du_dy, uv_cm, 'Họp Đảng ủy (so với số có mặt)')
    du_kdy = None if du_dy is None or uv_cm is None else uv_cm - du_dy
    ctx.update({
        'uv_tong': so(uv), 'uv_cm': so(uv_cm), 'uv_vm': so(uv_vm),
        'du_lydo_vang': v['du_lydo_vang'] or ('không.' if uv_vm == 0 else TRONG),
        'du_dy': so(du_dy), 'du_pt': pct(du_dy, uv), 'du_kdy': so(du_kdy), 'du_kpt': pct(du_kdy, uv),
        'du_lydo_khong': ly_do(du_kdy, v['du_lydo_khong']),
    })

    # Đoạn văn dài
    for k in NX_KEYS:
        v[k] = re.sub(r'\{\w+\}', TRONG, thay_the(v[k], v))
    ctx['nhan_thuc'] = lines(v['nhan_thuc']) or [TRONG * 6]
    ctx['ly_lich'] = lines(v['ly_lich']) or [TRONG * 6]
    uu, kd = lines(v['uu_diem']), lines(v['khuyet_diem'])
    gach = lambda xs: [x if x.startswith(('-', '+', '*')) else f'- {x}' for x in xs]
    ctx['uu_kd'] = ['* Ưu điểm:', *gach(uu or [TRONG * 6]), '* Khuyết điểm:', *gach(kd or [TRONG * 6])]
    return ctx


def ten_thu_muc(ho_ten):
    return re.sub(r'[\\/:*?"<>|]', '', ho_ten).strip()


def xuat_ho_so(data, thu_muc=None, to_do=True):
    """Điền 5 mẫu, lưu vào HoSo/<họ tên>/. to_do: tô đỏ chữ được điền. Trả về thư mục kết quả."""
    ctx = build_context(data)
    out = Path(thu_muc) if thu_muc else HOSO_DIR / ten_thu_muc(data['ho_ten'])
    out.mkdir(parents=True, exist_ok=True)
    for key, ten in MAU:
        tpl = DocxTemplate((MAU_DO_DIR if to_do else MAU_DIR) / f'{key}.docx')
        tpl.render(ctx, autoescape=True)
        tpl.save(out / f'{ten}.docx')
    luu_json(out / 'thong_tin.json', data)
    luu_json(DON_VI_FILE, {k: data.get(k, '') for k in DON_VI_KEYS})
    return out


def luu_json(path, data):
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')


def doc_excel(path):
    """Đọc file "Mau nhap thong tin.xlsx": cột "Mã" là khóa, cột "Nội dung nhập" là giá trị."""
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    for h, header in enumerate(rows):
        cells = [str(x or '').strip() for x in header]
        if 'Mã' in cells and 'Nội dung nhập' in cells:
            ck, cv = cells.index('Mã'), cells.index('Nội dung nhập')
            break
    else:
        raise LoiNhap('File Excel không đúng mẫu: thiếu cột "Mã" hoặc "Nội dung nhập".')
    data = {}
    for r in rows[h + 1:]:
        key = str(r[ck] or '').strip() if ck < len(r) else ''
        if key not in FIELDS:
            continue
        val = r[cv] if cv < len(r) else None
        if isinstance(val, (datetime.date, datetime.datetime)):
            val = val.strftime('%d/%m/%Y')
        elif isinstance(val, float) and val.is_integer():
            val = str(int(val))
        val = '' if val is None else str(val).strip()
        for m in MAU_NX.get(key, []):   # chọn tên mẫu nhận xét -> lấy nội dung mẫu
            if val == m['ten']:
                val = m['noi_dung']
        data[key] = val
    return data


def doc_json(path):
    try:
        return json.loads(Path(path).read_text(encoding='utf-8'))
    except (OSError, ValueError):
        return {}


# ---------------------------------------------------------------- giao diện
def chay_giao_dien():
    import tkinter as tk
    from tkinter import ttk, messagebox, filedialog

    root = tk.Tk()
    root.title('Lập hồ sơ kết nạp đảng viên')
    root.geometry('820x640')
    root.option_add('*Font', ('Segoe UI', 10))
    style = ttk.Style()
    style.configure('Ghi.TLabel', foreground='#777')

    widgets = {}
    nb = ttk.Notebook(root)
    nb.pack(fill='both', expand=True, padx=8, pady=(8, 0))

    for tab_name, fields in TABS:
        outer = ttk.Frame(nb)
        nb.add(outer, text=tab_name)
        canvas = tk.Canvas(outer, highlightthickness=0)
        sb = ttk.Scrollbar(outer, orient='vertical', command=canvas.yview)
        frame = ttk.Frame(canvas, padding=12)
        frame.bind('<Configure>', lambda e, c=canvas: c.configure(scrollregion=c.bbox('all')))
        win = canvas.create_window((0, 0), window=frame, anchor='nw')
        canvas.bind('<Configure>', lambda e, c=canvas, w=win: c.itemconfigure(w, width=e.width))
        canvas.configure(yscrollcommand=sb.set)
        canvas.pack(side='left', fill='both', expand=True)
        sb.pack(side='right', fill='y')
        canvas.bind('<Enter>', lambda e, c=canvas: c.bind_all(
            '<MouseWheel>', lambda ev: c.yview_scroll(int(-ev.delta / 120), 'units')))
        frame.columnconfigure(1, weight=1)

        row = 0
        for key, label, kind, default, note in fields:
            if kind == 'long':
                head = ttk.Frame(frame)
                head.grid(row=row, column=0, columnspan=2, sticky='ew', pady=(10, 2))
                ttk.Label(head, text=label).pack(side='left')
                t = tk.Text(frame, height=7, wrap='word', undo=True)
                t.grid(row=row + 1, column=0, columnspan=2, sticky='ew')
                widgets[key] = t
                presets = MAU_NX.get(key, [])
                if presets:
                    cb = ttk.Combobox(head, state='readonly', width=34, values=[x['ten'] for x in presets])
                    cb.current(0)

                    def dung(mode, key=key, cb=cb, t=t, presets=presets):
                        text = thay_the(presets[cb.current()]['noi_dung'], get_data())
                        if mode == 'them' and t.get('1.0', 'end').strip():
                            t.insert('end', '\n' + text)
                        else:
                            t.delete('1.0', 'end')
                            t.insert('1.0', text)
                    ttk.Button(head, text='Thêm vào', command=lambda d=dung: d('them')).pack(side='right')
                    ttk.Button(head, text='Thay', command=lambda d=dung: d('thay')).pack(side='right', padx=4)
                    cb.pack(side='right')
                    ttk.Label(head, text='Mẫu:', style='Ghi.TLabel').pack(side='right', padx=(10, 4))
                row += 2
                continue
            ttk.Label(frame, text=label).grid(row=row, column=0, sticky='w', pady=3, padx=(0, 10))
            e = ttk.Entry(frame, width=40 if kind == 'text' else 14)
            e.grid(row=row, column=1, sticky='ew' if kind == 'text' else 'w', pady=3)
            widgets[key] = e
            if note:
                ttk.Label(frame, text=note, style='Ghi.TLabel').grid(row=row + 1, column=1, sticky='w')
                row += 1
            row += 1
        if tab_name == 'Nhận xét':
            ttk.Label(frame, style='Ghi.TLabel', text='Chọn mẫu rồi bấm "Thay" (thay toàn bộ) hoặc "Thêm vào" (nối thêm). Mỗi dòng là một đoạn; sửa lại cho đúng từng người.'
                      ).grid(row=row, column=0, columnspan=2, sticky='w', pady=6)

    def get_data():
        return {k: (w.get('1.0', 'end').strip() if isinstance(w, tk.Text) else w.get().strip())
                for k, w in widgets.items()}

    def set_data(data):
        for k, w in widgets.items():
            val = str(data.get(k, ''))
            if isinstance(w, tk.Text):
                w.delete('1.0', 'end')
                w.insert('1.0', val)
            else:
                w.delete(0, 'end')
                w.insert(0, val)

    def ho_so_moi():
        data = defaults()
        data.update({k: v for k, v in doc_json(DON_VI_FILE).items() if k in DON_VI_KEYS})
        set_data(data)
        nb.select(1)
        widgets['ho_ten'].focus_set()

    def mo_ho_so():
        f = filedialog.askopenfilename(initialdir=HOSO_DIR if HOSO_DIR.exists() else HERE,
                                       title='Chọn thong_tin.json của hồ sơ',
                                       filetypes=[('Hồ sơ', 'thong_tin.json'), ('JSON', '*.json')])
        if f:
            data = defaults()
            data.update(doc_json(f))
            set_data(data)

    def nhap_excel():
        f = filedialog.askopenfilename(initialdir=HERE.parent, title='Chọn file Excel thông tin',
                                       filetypes=[('Excel', '*.xlsx')])
        if not f:
            return
        try:
            data = defaults()
            data.update(doc_excel(f))
        except LoiNhap as e:
            messagebox.showwarning('Không đọc được file', str(e))
            return
        except Exception as e:
            messagebox.showerror('Không đọc được file', f'Hãy đóng file trong Excel rồi thử lại.\n{e}')
            return
        set_data(data)
        nb.select(1)
        status.set(f'Đã nhập thông tin từ {Path(f).name}. Kiểm tra lại rồi bấm XUẤT.')

    def luu():
        data = get_data()
        if not data['ho_ten']:
            messagebox.showwarning('Thiếu thông tin', 'Chưa nhập họ tên quần chúng')
            return
        out = HOSO_DIR / ten_thu_muc(data['ho_ten'])
        out.mkdir(parents=True, exist_ok=True)
        luu_json(out / 'thong_tin.json', data)
        luu_json(DON_VI_FILE, {k: data[k] for k in DON_VI_KEYS})
        status.set(f'Đã lưu thông tin vào {out}')

    def xuat():
        try:
            out = xuat_ho_so(get_data(), to_do=to_do.get())
        except LoiNhap as e:
            messagebox.showwarning('Kiểm tra lại thông tin', str(e))
            return
        except PermissionError:
            messagebox.showerror('Không ghi được file', 'Hãy đóng các file Word của hồ sơ này đang mở rồi xuất lại.')
            return
        status.set(f'Đã xuất 5 mẫu vào {out}')
        if messagebox.askyesno('Xong', f'Đã xuất 5 mẫu vào:\n{out}\n\nMở thư mục?'):
            os.startfile(out)

    bar = ttk.Frame(root, padding=8)
    bar.pack(fill='x')
    ttk.Button(bar, text='Hồ sơ mới', command=ho_so_moi).pack(side='left')
    ttk.Button(bar, text='Nhập từ Excel...', command=nhap_excel).pack(side='left', padx=(6, 0))
    ttk.Button(bar, text='Mở hồ sơ...', command=mo_ho_so).pack(side='left', padx=6)
    ttk.Button(bar, text='Lưu thông tin', command=luu).pack(side='left')
    ttk.Button(bar, text='XUẤT 5 MẪU WORD', command=xuat).pack(side='right')
    to_do = tk.BooleanVar(value=True)
    ttk.Checkbutton(bar, text='Tô đỏ thông tin đã nhập', variable=to_do).pack(side='right', padx=10)
    status = tk.StringVar(value='Tab "Đơn vị" chỉ cần nhập một lần, lần sau tự điền lại.')
    ttk.Label(root, textvariable=status, style='Ghi.TLabel', padding=(10, 0, 10, 6)).pack(fill='x')

    ho_so_moi()
    nb.select(0 if not doc_json(DON_VI_FILE).get('truong') else 1)
    root.mainloop()


if __name__ == '__main__':
    if len(sys.argv) > 1:  # chạy không giao diện: python ket_nap.py thong_tin.json|.xlsx [thư mục]
        d = defaults()
        d.update(doc_excel(sys.argv[1]) if sys.argv[1].lower().endswith('.xlsx') else doc_json(sys.argv[1]))
        print('Đã xuất vào', xuat_ho_so(d, sys.argv[2] if len(sys.argv) > 2 else None))
    else:
        chay_giao_dien()
